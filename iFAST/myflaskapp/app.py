import io
import multiprocessing
import threading
from werkzeug.middleware.proxy_fix import ProxyFix
import time
from math import ceil
import json
from flask import Flask, request, render_template, redirect, session, flash, send_file, jsonify, Response, stream_with_context,url_for 
from openpyxl import Workbook
from Database import Database
import subprocess
import socket
import paramiko
import logging
import os
import tempfile
from paramiko import SSHClient, AutoAddPolicy, SSHException
import datetime
from flask_socketio import SocketIO
import pandas as pd
from io import StringIO
import random
from flask_cors import CORS 


app = Flask(__name__, static_folder='static')
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_port=1)
logging.basicConfig(level=logging.DEBUG, format='%(asctime)s %(levelname)s: %(message)s', filename='debug.log')
info_handler = logging.FileHandler('info.log')
info_handler.setLevel(logging.INFO)
info_handler.setFormatter(logging.Formatter('%(asctime)s %(message)s'))
app.logger.addHandler(info_handler)

app.config['DEBUG'] = True
app.secret_key = 'asdsdfsdfs13sdf_df%&'
app.config['MYSQL_HOST'] = 'localhost'
app.config['MYSQL_USER'] = 'root'
app.config['MYSQL_PASSWORD'] = 'Tools@123'
app.config['MYSQL_DB'] = 'iFAST'
app.config['SSH_DETAILS'] = {
    'host': None,
    'user': None,
    'password': None
}
app.config['SSH_COMMAND'] = None

socketio = SocketIO(app, async_mode='gevent', cors_allowed_origins="*")
CORS(app, resources={r"/*": {"origins": "*"}})

db = Database(app)
logs = []
processes = []
current_script_args_list = []
@app.route("/")
def index():
    return render_template("new_index.html")

@app.route("/test_redirect")
def test_redirect():
    return render_template('newpage.html')

@app.route("/login", methods=['GET', 'POST'])
def login():
    inactive_modal = request.args.get('inactive_modal', "false").lower()
    return render_template('new_index.html', inactive_modal=inactive_modal)

@app.route("/login_succ", methods=['POST', 'GET'])
def login_succ():
    if request.method == "POST":
        username = request.form.get('username')
        passwd = request.form.get('passwd')
        mode = request.form.get('mode')
        session['mode'] = mode
        query = 'SELECT * FROM user WHERE knox_id=%s AND password=MD5(%s)'
        values = (username, passwd)
        acct = db.execute_query(query, values)
        if acct:
            acct = dict(acct[0])
            session['user_name'] = username
            app.logger.info(f"Account details: {acct}")
            if acct['status'] == 'active':
                session['user_name'] = username
                session['role'] = acct
                app.logger.info(f"Account details found redirecting")
                redirect_url = url_for('newpage', _external=True, _scheme='http')
                redirect_url = redirect_url.replace('http://105.57.243.148', 'http://105.57.243.148:12258')
                app.logger.info(f"Redirecting to {redirect_url}")
                response = redirect(redirect_url)
                app.logger.info(f"Response status code: {response.status_code}")
                return response
            else:
                return redirect(url_for('login', inactive_modal="true"))
        else:
            flash('Wrong Creds. If you do not have an account, please apply through JIRA.')
            return redirect(url_for('login'))
    elif request.method == "GET":
        return "Method not allowed", 405

@app.route("/newpage",methods=['GET','POST'])
def newpage():
    app.logger.info('Request to /newpage')
    query = "SELECT * FROM known_knowledge_base"
    query2 = 'select * from user'
    roles = ['admin','user']
    status = ['active','inactive']
    acct = db.execute_query(query)
    user_tuples = db.execute_query(query2)
    users = list(user_tuples)
    total_entries = len(users)
    entries_per_page = 5  # Assuming 5 entries per page
    total_pages = ceil(total_entries / entries_per_page)
    page_number = int(request.args.get('page', 1))
    res = db.execute_query(query, fetch_all=True)
    import pdb
    pdb.set_trace()
    return render_template("ifast.html",res = res,roles=roles,users=users,
                           total_entries=total_entries,
                           total_pages=total_pages,
                           page_number=page_number,
                           status=status)

@app.route('/landing_page', methods=['GET', 'POST'])
def landing_page():
    return render_template("ifast_main.html")

@app.route("/action_exe/<issue>", methods=['POST', 'GET'])
def action_exe(issue):
    sequence_num = 1
    app.logger.info(f'SELECT 1 FROM scripts WHERE issue_id = {issue} AND sequence_num = {sequence_num}')
    query = None
    if db.execute_query("SELECT 1 FROM scripts WHERE issue_id = %s AND sequence_num = %s", (issue, sequence_num), fetch_all=False):
        query = """
            SELECT kb.*, s.*
            FROM known_knowledge_base kb
            INNER JOIN scripts s ON kb.issue = s.issue_id AND s.sequence_num = 1
            WHERE kb.issue = %s
        """
    elif db.execute_query("SELECT 1 FROM get_file WHERE issue_id = %s AND sequence_num = %s", (issue, sequence_num), fetch_all=False):
        query = """
            SELECT kb.*, g.*
            FROM known_knowledge_base kb
            INNER JOIN get_file g ON kb.issue = g.issue_id AND g.sequence_num = 1
            WHERE kb.issue = %s
        """
    elif db.execute_query("SELECT 1 FROM grep WHERE issue_id = %s AND sequence_num = %s", (issue, sequence_num), fetch_all=False):
        query = """
            SELECT kb.*, gr.*
            FROM known_knowledge_base kb
            INNER JOIN grep gr ON kb.issue = gr.issue_id AND gr.sequence_num = 1
            WHERE kb.issue = %s
        """
    res = db.execute_query(query, (issue,))
    if res:
        res = res[0]
    return render_template('action_exe.html', issue=issue, record=res)

@app.route('/send_message', methods=['POST', 'GET'])
def send_message():
    issue_id = request.form.get('issue_id')
    SW = request.form.get('sw')
    issue_st =request.form.get('issue_st')
    issuecat =request.form.get('issuecat')
    POC     =request.form.get('POC')
    sequence_id=request.form.get('sequence_id')
    action_type=request.form.get('action_type')
    uid      =request.form.get('uid')
    pwd      =request.form.get('pwd')
    ip       =request.form.get('ip')
    proceed  =request.form.get('proceed')
    counter_info =request.form.get('counter_info')
    check_condition=request.form.get('check_condition')
    email    =request.form.get('email')
    regex    =request.form.get('regex')   
    query = "SELECT total_steps FROM known_knowledge_base WHERE issue=%s"
    values = (issue_id,)
    cnt = db.execute_query(query, values)
    app.logger.info(f"Action: {issue_st} {issuecat} {POC} {sequence_id} {action_type} {uid} {pwd} {ip} {proceed} {counter_info} {check_condition} {email} {regex}")
    app.logger.info(f"{request.form}")

    app.config['SSH_DETAILS'] = {
        'host': '10.100.1.241',
        'user': 'toolsuser',
        'password': 'Tools@123'
    }

    app.config['SSH_COMMAND'] = 'bash /opt/vz_raw_data/Scripts/Test/iFAST_OM_Executer/om_exec.sh "NR_PRACH_BEAM| RachPreambleAPerBeam,RachPreambleACFRAPerBeam,NumMSG3PerBeam,NumMSG3CFRAPerBeam\nCU-CSL| CauseInternalRandomAccessProblem" "RachPreambleAPerBeam+RachPreambleACFRAPerBeam,=,0\nNumMSG3PerBeam+NumMSG3CFRAPerBeam,=,0\nCauseInternalRandomAccessProblem,!=,0\n#proceed,No" "24593002050_5GDU_CATAWBA_SOUTH-CLMB,24593002373_5GDU_MERCHANTS_LANDING-CLMB,24593003029_5GDU_FOSTORIA-CLMB" "21_1"'

    return jsonify({'message': 'Execution started'}), 200

@app.route('/get_updates', methods=['GET'])
def get_updates():
    def generate():
        ssh_details = app.config.get('SSH_DETAILS')
        ssh_command = app.config.get('SSH_COMMAND')

        app.logger.info(f"SSH_DETAILS during get_updates: {ssh_details}")
        app.logger.info(f"SSH_COMMAND during get_updates: {ssh_command}")

        if not ssh_details['host'] or not ssh_details['user'] or not ssh_details['password'] or not ssh_command:
            app.logger.error('Missing SSH details or command')
            yield "data: Error: Missing SSH details or command\n\n"
            return

        ssh_host = ssh_details['host']
        ssh_user = ssh_details['user']
        ssh_password = ssh_details['password']

        app.logger.info('Starting SSH connection for updates')
        try:
            ssh_client = paramiko.SSHClient()
            ssh_client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
            ssh_client.connect(ssh_host, username=ssh_user, password=ssh_password)

            stdin, stdout, stderr = ssh_client.exec_command(ssh_command)

            for line in iter(stdout.readline, ''):
                app.logger.debug(f'Output: {line.strip()}')
                yield f"data: {line}\n\n"

            ssh_client.close()
            app.logger.info('SSH connection closed')
        except Exception as e:
            app.logger.error(f'Error in SSH connection: {str(e)}')
            yield f"data: Error: {str(e)}\n\n"

    return Response(stream_with_context(generate()), content_type='text/event-stream')

@app.route('/test-db')
def test_db():
    if db.check_connection():
        app.logger.info('Database connection successful!')
        return jsonify({'message': 'Database connection successful!'})
    else:
        app.logger.error('Failed to connect to database!')
        return jsonify({'message': 'Failed to connect to database!'})

@app.route('/test-socket')
def test_socket():
    try:
        server_address = ('localhost', 8001)
        client_socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        client_socket.connect(server_address)
        message = "Hello, server!"
        client_socket.sendall(message.encode())
        server_response = client_socket.recv(1024).decode()
        client_socket.close()
        return jsonify({'message': 'Socket connection successful!', 'server_response': server_response})
    except Exception as e:
        app.logger.error(f'Error testing socket: {str(e)}')
        return jsonify({'error': str(e)})

@app.route("/fetch_kb", methods=['POST','GET'])
def fetch_kb():
    role = None
    if 'role' in session:
        role = session['role']
    page = int(request.args.get("page", 1))
    per_page = 1
    total_records_query = "SELECT COUNT(*) as cnt FROM known_knowledge_base"
    total_records = db.execute_query(total_records_query, fetch_all=True)[0]['cnt']
    query = "SELECT * FROM known_knowledge_base"
    res = db.execute_query(query, fetch_all=True)
    # Paginate the records
    offset = (page - 1) * per_page
    # pagination_users = res[offset: offset + per_page]
    status_info = [user['Status'] for user in res]
    return render_template('viewknowledgebase.html', users=res ,status_info=status_info, roles=role)

@app.route("/fetch_issues", methods=['POST', 'GET'])
def fetch_issues():
    if request.method == 'POST':
        search_value = request.form.get('searchValue', '')
        select_SW = request.form.get('select_SW', '')
        select_Tech = request.form.get('select_Tech', '')
        select_kpi = request.form.get('select_kpi', '')
        select_Area = request.form.get('select_Area', '')

        query = "SELECT * FROM known_knowledge_base WHERE 1=1"
        if search_value:
            query += f" AND issue_category LIKE '%{search_value}%'"
        if select_SW and select_SW != 'nothing_sw':
            query += f" AND SW = '{select_SW}'"
        if select_Tech and select_Tech != 'nothing_tech':
            query += f" AND Technology = '{select_Tech}'"
        if select_kpi and select_kpi != 'nothing_kpi':
            query += f" AND Major_KPI_Degradation = '{select_kpi}'"
        if select_Area and select_Area != 'nothing_area':
            query += f" AND Area = '{select_Area}'"

        res = db.execute_query(query, fetch_all=True)
        return jsonify(res)
    else:  # GET request (initial page load)
        return render_template('viewOverallIssues.html', users=[])


@app.route('/execute_all')
def execute_all():
    return render_template('execute_all.html')

@app.route("/logout", methods=['POST', 'GET'])
def logout():
    app.logger.info("Logging out user")
    session.clear()
    redirect_url = f"http://105.57.243.148:12258/"
    app.logger.info(f"Redirecting to: {redirect_url}")
    return redirect(redirect_url)

@app.route("/kb_hit", methods=['POST', 'GET'])
def kb_hit():
    return render_template('kb_GUI.html')

@app.route("/act_srch/<issue>", methods=['POST', 'GET'])
def act_srch(issue):
    query = "SELECT * FROM known_knowledge_base k,scripts s WHERE k.issue = %s and k.issue = s.issue_id"
    res = db.execute_query(query, (issue,),fetch_all=True)
    return render_template('viewSingleIssue.html', users=res)


@app.route("/testpoc", methods=['POST', 'GET'])
def testpoc():
   return render_template("testexec.html")

def run_script(script_args, output_queue):
    try:
        output_queue.put("[START OF PROCESS]")

        script = ['python3', '-u', '/home/sieluser/myflaskapp/scripts/script_exec.py']
        output_queue.put(f"Executing script with arguments: {script + script_args}")

        process = subprocess.Popen(script + script_args,
                                   stdout=subprocess.PIPE,
                                   stderr=subprocess.PIPE,
                                   universal_newlines=True,
                                   bufsize=1)
        output_queue.put("[START OF SCRIPT]")

        for stdout_line in iter(process.stdout.readline, ''):
            output_queue.put(f"OUTPUT: {stdout_line.strip()}")

        process.stdout.close()
        process.wait()

        for stderr_line in iter(process.stderr.readline, ''):
            output_queue.put(f"ERROR: {stderr_line.strip()}")

        process.stderr.close()

        if process.returncode != 0:
            output_queue.put(f"ERROR: Process failed with return code {process.returncode}.")

        output_queue.put("[END OF PROCESS]")

    except Exception as e:
        output_queue.put(f"ERROR: {str(e)}")


def run_scripts(process_id, script_args, output_queue):
    app.logger.info(f"Process {process_id}: Called run scripts with args {script_args}")
    try:
        script_path = '/home/sieluser/myflaskapp/scripts/ifast_executer.py'
        scripts = [(['python3', '-u', script_path], script_args)]
        app.logger.info(f"[START OF PROCESS {process_id}]")
        output_queue.put(f"[START OF PROCESS {process_id}]")

        pid = os.getpid()
        for script, args in scripts:
            try:
                app.logger.info(f"Process {process_id}: Preparing to execute script {script} with args {args}")
                process = subprocess.Popen(script + args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True, bufsize=1)

                app.logger.info(f"Process {process_id}: Subprocess started with PID {process.pid}")

                for stdout_line in iter(process.stdout.readline, ''):
                    app.logger.info(f"Process {process_id} stdout: {stdout_line.strip()}")
                    output_queue.put(f"Process {process_id} OUTPUT: {stdout_line.strip()}")

                process.stdout.close()
                process.wait()

                for stderr_line in iter(process.stderr.readline, ''):
                    app.logger.error(f"Process {process_id} stderr: {stderr_line.strip()}")
                    output_queue.put(f"Process {process_id} ERROR: [ERROR] {stderr_line.strip()}")

                process.stderr.close()

                if process.returncode != 0:
                    app.logger.error(f"Process {process_id}: Subprocess exited with return code {process.returncode}")
                    output_queue.put(f"Process {process_id} ERROR: [ERROR OCCURRED with {script[2]}]")

            except Exception as e:
                app.logger.error(f"Process {process_id}: Exception occurred: {str(e)}")
                output_queue.put(f"Process {process_id}: [ERROR OCCURRED with {script[2]}: {str(e)}]")

        output_queue.put(f"[END OF PROCESS {process_id}]")

    except Exception as e:
        app.logger.error(f"Process {process_id}: Unhandled exception: {str(e)}")
        output_queue.put(f"[ERROR OCCURRED in Process {process_id}: {str(e)}]")

def execute_script():
    app.logger.info("Starting parallel script execution")

    num_pools = 5
    queues = [multiprocessing.Queue() for _ in range(num_pools)]
    processes = []

    for pool_index in range(num_pools):
        process_id = f"{pool_index+1}"
        p = multiprocessing.Process(target=run_scripts, args=(process_id, queues[pool_index]))
        processes.append(p)
        p.start()

    def generate():
        try:
            for pool_index in range(num_pools):
                while not queues[pool_index].empty():
                    message = queues[pool_index].get()
                    yield f"data: {message}\n\n"

                processes[pool_index].join()

            yield "data: [END OF STREAM]\n\n"

        except Exception as e:
            app.logger.error(f"Error occurred: {str(e)}")
            yield "data: [ERROR OCCURRED]\n\n"
        finally:
            app.logger.info("Parallel script execution completed.")

    return Response(stream_with_context(generate()), mimetype='text/event-stream')

@app.route('/execute')
def execute():
    return execute_script()

# @app.route("/save_singleissue", methods=['POST', 'GET'])
# def save_singleissue():
#     issue = request.form.get('msId')
#     poc = request.form.get('POC')
#     sw = request.form.get('swname')
#     status = request.form.get('status')
#     area = request.form.get('area')
#     tech = request.form.get('tech')
#     visible = request.form.get('Visibility')
#     kpi = request.form.get('kpi')
#     detect = request.form.get('detect')
#     cell = request.form.get('cell')
#     desc = request.form.get('desc')
#     rca = request.form.get('rca')
#     impact = request.form.get('impact')
#     script = request.form.get('Detection_Methodology_Log_Script')
#     oms = request.form.get('Detection_Methodology_OMs')
#     wa = request.form.get('Interim_Solution_Workaround')
#     pf = request.form.get('Permanent_Fix')
#     ic = request.form.get('ic')
#     print("status",status)
#     update_query = "UPDATE known_knowledge_base SET Poc=%s, SW=%s, Status=%s, Area=%s, Technology=%s, Major_KPI_Degradation=%s, Issue_Category=%s, " \
#                    "Description=%s, Root_Cause=%s, Impact=%s, Detection_Methodology_OMs=%s, Detection_Methodology_Log_Script=%s, " \
#                    "Interim_Solution_Workaround=%s,Visibility=%s, Permanent_Fix=%s, Issue_Detection_Category=%s, Cell_or_Site_Level=%s WHERE issue=%s"
#     values = (poc, sw, status, area, tech, kpi, ic, desc, rca, impact, oms, script, wa, visible, pf, detect, cell, issue)
#     try:
#         db.execute_query(update_query, values, commit=True)
#         return jsonify({'message': "updated succesfully!! Click on next to update action steps"})
#     except Exception as e:
#         return jsonify({'error': f"Something went wrong {e}"})

@app.route("/insert_kb",methods=['GET','POST'])
def insert_kb():
    app.logger.info("INFODODO")
    SW_val = request.form.get('issue_sw')
    status = request.form.get('status')
    area_val = request.form.get('area_val')
    tech = request.form.get('tech')
    Visibility = request.form.get('Visibility')
    kpi = request.form.get('kpi')
    detect = request.form.get('detect')
    cell = request.form.get('cell')
    icat = request.form.get('icat')
    desc = request.form.get('desc')
    rca = request.form.get('rca')
    impact = request.form.get('impact')
    m_oms = request.form.get('m_oms')
    script = request.form.get('script')
    wa = request.form.get('wa')
    pf = request.form.get('pf')
    poc = request.form.get('POC')
    sel_query = "select * from known_knowledge_base where SW=%s and area=%s and Technology=%s" \
                "and Major_KPI_Degradation=%s"
    values = (SW_val,area_val,tech,kpi)
    try:
        app.logger.info(f"TETSTSTST")
        res = db.execute_query(sel_query,values)
        print(res)
        if res:
            app.logger.info("Executing query inside if")
            redirect_url = url_for('act_srch', issue=res['issue'])
            return jsonify({"message": f"Issue {res['issue']} already exists. Redirecting...","redirect_url": redirect_url})
        else:
            app.logger.info("Executing query inside else")
            insert_query="insert into known_knowledge_base(SW,Status,Area,Technology,Visibility,Major_KPI_Degradation," \
                         "Issue_Category,Description,Root_Cause,Impact,Detection_Methodology_OMs," \
                         "Detection_Methodology_Log_Script,Interim_Solution_Workaround,Permanent_Fix,poc,Issue_Detection_Category,Cell_or_Site_Level) values" \
                         "(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
            values=(SW_val,status,area_val,tech,Visibility,kpi,icat,desc,rca,impact,m_oms,script,wa,pf,poc,detect,cell)
            db.execute_query(insert_query, values, commit=True)
            app.logger.info("inserted")
            result = db.execute_query("SELECT LAST_INSERT_ID() as issue_id")
            app.logger.info(f"{result}")
            issue_id = result[0]['issue_id'] if result else None
            print(issue_id)
            if issue_id:
                print(issue_id)
                return jsonify({"message": "Inserted successfully", "issue_id": issue_id, "sw_val":SW_val, "status_val":status})
            else:
                return jsonify({"message": "Insertion failed"}), 500
    except Exception as e:
        app.logger.info(e)
        return {"error":f"{e}"}


@app.route("/lookup_onload", methods=['POST', 'GET'])
def lookup_onload():
    cursor = db.get_cursor()
    # Context variables to send to template
    cat_qry = "SELECT DISTINCT Issue_Category FROM known_knowledge_base "
    cursor.execute(cat_qry)
    categories = [i['Issue_Category'] for i in cursor.fetchall()]
    sw_qry = "SELECT DISTINCT SW FROM known_knowledge_base "
    cursor.execute(sw_qry)
    softwares = [i['SW'] for i in cursor.fetchall()]
    tech_qry = "SELECT DISTINCT Technology FROM known_knowledge_base"
    cursor.execute(tech_qry)
    technologies = [i['Technology'] for i in cursor.fetchall()]
    kpi_qry = "SELECT DISTINCT Major_KPI_Degradation FROM known_knowledge_base "
    cursor.execute(kpi_qry)
    kpis = [i['Major_KPI_Degradation'] for i in cursor.fetchall()]
    area_qry = "SELECT DISTINCT AREA FROM known_knowledge_base "
    cursor.execute(area_qry)
    areas = [i['AREA'] for i in cursor.fetchall()]
    context = {
        'categories': categories,
        'softwares': softwares,
        'technologies': technologies,
        'kpis': kpis,
        'areas': areas,
    }
    if request.method == "POST":
        filter_attr = []
        category = request.form.get('issue_cat')
        software = request.form.get('select_SW', '')
        technology = request.form.get('select_Tech', '')
        kpi = request.form.get('select_kpi', '')
        area = request.form.get('select_Area', '')

        # Validate required fields
        if not category or not technology:
            msg = 'Missing required fields'
            flash(msg, 'danger')
            return render_template('lookup_onload.html', **context)

        if len(software) > 0:
            filter_attr.append({'SW': software})
        if len(kpi) > 0:
            filter_attr.append({'Major_KPI_Degradation': kpi})
        if len(area) > 0:
            filter_attr.append({'Area': area})

        context['category_val'] = category
        context['software_val'] = software
        context['technology_val'] = technology
        context['kpi_val'] = kpi
        context['area_val'] = area

        unique_keys = [key for d in filter_attr for key in d]
        unique_values = [f'%{value}%' for d in filter_attr for value in d.values()]

        # Construct the query dynamically based on the number of categories
        #category_placeholders = ','.join(['%s'] * len(category))
        filters = ' AND '.join(f"{key} LIKE %s" for key in unique_keys) if unique_keys else ''
        final_query = f'SELECT * FROM known_knowledge_base WHERE Issue_Category like %s AND Technology LIKE %s' + (' AND ' + filters if filters else '')

        # Prepare parameters for the query
        params = [f'%{category}%'] + [f'%{technology}%'] + unique_values
        print(params)
        # Execute the query
        try:
            cursor.execute(final_query, params)
            result = cursor.fetchall()
            context['results'] = list(result)
        except MySQLdb.ProgrammingError as e:
            print(f"Database error: {e}")
            flash('An error occurred while querying the database.', 'danger')

        return render_template('lookup_onload.html', **context)

    elif request.method == "GET":
        return render_template('lookup_onload.html', **context)

    else:
        response = {'message': f'Invalid Method, Method {request.method} Not Supported'}
        flash(response['message'], 'danger')
        return response



@app.route("/bulk_upload", methods=['POST', 'GET'])
def bulk_upload():
    if request.method == "POST":
        uploaded_file = request.files['file']
        extension = uploaded_file.filename.split(".")[1]

        if uploaded_file.filename != '' and (extension == "csv" or extension == "xlsx" or extension == "xls" or extension == "tsv"):
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], uploaded_file.filename)
            uploaded_file.save(file_path)
            try:
                df = None
                if extension == "csv":
                    df = pd.read_csv(file_path)
                elif extension == "tsv":
                    df = pd.read_csv(file_path, delimiter="\t")
                else:
                    df = pd.read_excel(file_path, index_col=None, engine='openpyxl')

                total_records = len(df)
                error_records = parsecsv(extension, file_path)
            except Exception as e:
                print(f"Error reading Excel file: {str(e)}")
                return "Error reading Excel file"

            def generate():
                for i, (index, row) in enumerate(df.iterrows()):
                    progress = (i + 1) / total_records * 100
                    if i < len(error_records):
                        error_message = error_records.get['error_message']
                    else:
                        error_message = "No error message"
                    yield f"data: {progress:.2f}\n\n"

            response = Response(generate(), content_type="text/event-stream")
            response.headers["Cache-Control"] = "no-cache"
            # return error_records
            flash(error_records['message'], 'info')
            return redirect(url_for('fetch_kb'))
        else:
            return "Failure"

    return redirect(url_for('kb'))



def parsecsv(extension,csvfile):
    error_records = []
    if extension == "csv":
        df = pd.read_csv(csvfile)
    elif extension == "tsv":
        df = pd.read_csv(csvfile, delimiter="\t")
    else:
        df = pd.read_excel(csvfile, index_col=None, engine='openpyxl')

    cursor = db.get_cursor()
    columns = df.columns.difference(["Issue#", "Issue seen S/W ", "Status", "Area", "Technology", "Visibility", "Major KPI Degradation", "Issue Category/Issue Name", "Description",
                                     "Root Cause", "Impact", "Detection Methodology Oms", "Detection Methodology Log/Script", "Interim Solution/Workaround", "Permanent Fix", "PoC",
                                    "Issue Detection Category", "Cell / Site Level"])
    df[columns] = df[columns].fillna("NA")
    cursor.execute("SHOW COLUMNS FROM known_knowledge_base")
    all_cols = [row['Field'] for row in cursor.fetchall()]
    col_ex_issue = [i for i in all_cols if i not in ['issue', 'Issue_Name', 'total_steps', 'email_ids']]
    col_ex_issue[18], col_ex_issue[19] = col_ex_issue[19], col_ex_issue[18]
    qry1 = 'SELECT {} FROM KNOWN_KNOWLEDGE_BASE WHERE '
    dupe_qry = qry1 +' AND '.join(x for x in ['{}=%s'.format(x) for x in col_ex_issue])
    dupe_qry = dupe_qry.format(','.join(str(x) for x in all_cols))
    import re
    pattern = r'#\w+(?=[,=])'

    def enclose_with_backticks(match):
        return f'`{match.group(0)}`'

    final_qry = re.sub(pattern, enclose_with_backticks, dupe_qry)
    duplicate_ids = []
    for i, row in df.iterrows():
        try:
            ls1 = [row[x] if row[x] is not None else '' for x in range(1,15)]
            ls2 = [row[y] if row[y] is not None else '' for y in range(17,31)]
            ls1.extend(ls2)
            cursor.execute(final_qry, tuple(ls1))
        except Exception as  e:
            print(e)
        result = cursor.fetchone()
        if result:
            duplicate_ids.append(result['issue'])
            continue
        # code block to check for existing issue
        existing_issue_ids = set()
        existing_query = 'SELECT issue from known_knowledge_base'
        cursor.execute(existing_query)
        existing_ids = cursor.fetchall()
        if existing_ids:
            existing_issue_ids.update(record['issue'] for record in  existing_ids)
        try:
            issue = int(row['Issue#'])
            sw = row['Issue seen S/W ']
            status = row['Status']
            area = row['Area']
            technology = row['Technology']
            visibility = row['Visibility']
            kpi_degradaton = row['Major KPI Degradation']
            category = row['Issue Category/Issue Name']
            description = row['Description']
            cause = row['Root Cause']
            impact = row['Impact']
            oms = row['Detection Methodology Oms']
            log_script = row['Detection Methodology Log/Script']
            solution = row['Interim Solution/Workaround']
            fix = row['Permanent Fix']
            poc = row['PoC']
            detection = row['Issue Detection Category']
            level = row['Cell / Site Level']
            # periodicity =row['Periodicity']
            # family =row['#1 OM FAMILY and OM Counters']
            # step = row['OM_Detection_Step']
            # check = row['#1 OM validation check']
            # cli = row['CLI_Detection_Step']
            # cli_2 = row['#2 CLI']
            # grep = row['GREP_Condition']
            # grep_3 = row['#3 GREP']
            # file = row['GET_File']
            # file_4 = row['#4 GET_File']
            # actions = row['Sequence for configured actions']

            periodicity = row[20]
            family = row[22]
            step = row[21]
            check = row[23]
            cli = row[24]
            cli_2 = row[25]
            grep = row[26]
            grep_3 = row[27]
            file = row[28]
            file_4 = row[29]
            actions = row[30]

            if pd.isna(issue):
                print(f"record {i} : Req issue id is missing (1048, \"column 'issue' cannot be null\")")
                continue

            if pd.isna(sw):
                print(f"record {i} : Req Issue seen S/W is missing (1048, \"column 'Issue seen S/W' cannot be null\")")
                continue

            if pd.isna(status):
                print(f"record {i} : Req id Status missing (1048, \"column 'Status' cannot be null\")")
                continue

            if pd.isna(area):
                print(f"record {i} : Req Area is missing (1048, \"column 'Area' cannot be null\")")
                continue

            if pd.isna(technology):
                print(f"record {i} : Req Technology is missing (1048, \"column 'Technology' cannot be null\")")
                continue

            if pd.isna(visibility):
                print(f"record {i} : Req Visibility is missing (1048, \"column 'Visibility' cannot be null\")")
                continue

            if pd.isna(kpi_degradaton):
                print(f"record {i} : Req Major KPI Degradation is missing (1048, \"column 'Major KPI Degradation' cannot be null\")")
                continue


            if pd.isna(category):
                print(f"record {i} : Req Issue Category/Issue Name is missing (1048, \"column 'Issue Category/Issue Name' cannot be null\")")
                continue

            if pd.isna(description):
                print(f"record {i} : Req Description is missing (1048, \"column 'Description' cannot be null\")")
                continue

            if pd.isna(cause):
                print(f"record {i} : Req Root Cause is missing (1048, \"column 'Root Cause' cannot be null\")")
                continue

            if pd.isna(impact):
                print(f"record {i} : Req Impact is missing (1048, \"column 'Impact' cannot be null\")")
                continue

            if pd.isna(oms):
                print(f"record {i} : Req Detection Methodology Oms is missing (1048, \"column 'Detection Methodology Oms' cannot be null\")")
                continue

            if pd.isna(log_script):
                print(f"record {i} : Req Detection Methodology Log/Script is missing (1048, \"column 'Detection Methodology Log/Script' cannot be null\")")
                continue

            if pd.isna(solution):
                print(f"record {i} : Req Interim Solution/Workaround is missing (1048, \"column 'Interim Solution/Workaround' cannot be null\")")
                continue

            if pd.isna(fix):
                print(f"record {i} : Req Permanent Fix is missing (1048, \"column 'Permanent Fix' cannot be null\")")
                continue

            if pd.isna(poc):
                print(f"record {i} : Req PoC is missing (1048, \"column 'PoC' cannot be null\")")
                continue

            if pd.isna(detection):
                print(f"record {i} : Req Issue Detection Category is missing (1048, \"column 'Issue Detection Category' cannot be null\")")
                continue

            if pd.isna(level):
                print(f"record {i} : Req Cell/Site Level is missing (1048, \"column 'Cell/Site Level' cannot be null\")")
                continue


            if issue in existing_issue_ids:
                values = [
                          sw,
                          status,
                          area,
                          technology,
                          visibility,
                          kpi_degradaton,
                          category,
                          description,
                          cause,
                          impact,
                          oms,
                          log_script,
                          solution,
                          fix,
                          poc,
                          category,
                          detection,
                          level,
                          periodicity,
                          family,
                          step,
                          check,
                          cli,
                          cli_2,
                          grep,
                          grep_3,
                          file,
                          file_4,
                          actions,
                          issue
                          ]

                values = [value if not pd.isna(value) else '' for value in values]
                error_message = f"{issue} already present in DB!! Updated the record as uploaded thru Excel"
                update_qry = """UPDATE known_knowledge_base
                               SET SW=%s, Status=%s, Area=%s,
                               Technology=%s, Visibility=%s, Major_KPI_Degradation=%s,
                               Issue_Category=%s, Description=%s, Root_Cause=%s,
                               Impact=%s, Detection_Methodology_OMs=%s, Detection_Methodology_Log_Script=%s,
                               Interim_Solution_Workaround=%s, Permanent_Fix=%s, PoC=%s,
                               Issue_Name=%s, Issue_Detection_Category=%s, Cell_or_Site_Level=%s, Periodicity=%s,
                                OM_Family_OM_Counters=%s, OM_Detection_Step_Summary=%s, `#1_OM_validation_check`=%s,
                               CLI_Detection_Step_Summary=%s, `#2_CLI_RunScript_Check`=%s, Grep_Condition_Summary=%s,
                               `#3_GREP_Condition_Check`=%s, GET_File_Summary=%s, `#4_GETFile_Condition_Check`=%s,
                               Sequence_of_configuration=%s,  WHERE issue=%s """

                final_update_qry = re.sub(r'\s+', ' ',update_qry.replace('/n', '  '))
                cursor.execute(final_update_qry,  tuple(values))
                cursor.commit()
                cursor.close()
                # db.execute_query(, params= tuple(values), commit=True)

                error_records.append({'row_index': i, 'error_message': error_message})





                continue

            else:
                qry = 'SELECT issue FROM known_knowledge_base ORDER BY issue DESC LIMIT 1'
                cursor.execute(qry)
                last_issue_idx = cursor.fetchone()
                cur_idx = last_issue_idx["issue"] + 1

                values = [cur_idx,
                    sw,
                    status,
                    area,
                    technology,
                    visibility,
                    kpi_degradaton,
                    category,
                    description,
                    cause,
                    impact,
                    oms,
                    log_script,
                    solution,
                    fix,
                    poc,
                    category,
                    detection,
                    level,
                    periodicity,
                    family,
                    step,
                    check,
                    cli,
                    cli_2,
                    grep,
                    grep_3,
                    file,
                    file_4,
                    actions
                ]

                values = [value if not pd.isna(value) else '' for value in values]
                query = """INSERT INTO known_knowledge_base(issue, SW, Status, Area, Technology, Visibility, 
                Major_KPI_Degradation, Issue_Category, Description, Root_Cause, Impact, Detection_Methodology_OMs, 
                Detection_Methodology_Log_Script, Interim_Solution_Workaround, Permanent_Fix, PoC, Issue_Name,
                Issue_Detection_Category, Cell_or_Site_Level, Periodicity, OM_Family_OM_Counters, 
                OM_Detection_Step_Summary, `#1_OM_validation_check`, CLI_Detection_Step_Summary, 
                `#2_CLI_RunScript_Check`, Grep_Condition_Summary, `#3_GREP_Condition_Check`, GET_File_Summary, 
                `#4_GETFile_Condition_Check`, Sequence_of_configuration)VALUES(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);"""
                db.execute_query(query, params=values, commit=True)



        except Exception as e:
            error_message = f" was not inserted because of error:{e}"
            print(f"Error occured in row {i}: {str(e)}")
            error_records.append({'row_index': i, 'error_message': error_message})
            response = {'message': error_records}
            return response

    message = "File Uploaded Successfully!"
    response = {'message': message}
    return response

@app.route("/execute_selected", methods=['POST'])
def execute_selected():
    if request.method == "POST":
        context = {}
        category = str(request.form.get('issue_cat'))
        software = str(request.form.get('select_SW'))
        technology = str(request.form.get('select_Tech'))
        kpi = str(request.form.get('select_kpi'))
        area = str(request.form.get('select_Area'))
        regex = str(request.form.get('regex'))
        email_ids = str(request.form.get('email_id'))
        issues = [i for i in str(request.form.get('selected_issue_ids')).split(',')]
        context['category_val'] = category
        context['software_val'] = software
        context['technology_val'] = technology
        context['kpi_val'] = kpi
        context['area_val'] = area
        context['regex'] = regex
        context['email_ids'] = email_ids
        context['issues'] = issues
        cursor = db.get_cursor()
        placeholders = ','.join(['%s'] * len(issues))
        # result_qry = 'SELECT * FROM known_knowledge_base WHERE Issue_Category = %s AND SW = %s  AND Technology = %s AND Major_KPI_Degradation = %s AND Area = %s'
        result_qry = f'SELECT * FROM known_knowledge_base WHERE issue IN ({placeholders})'
        cursor.execute(result_qry, tuple(issues))
        # cursor.execute(result_qry, (category, software, technology, kpi, area ))
        result = cursor.fetchall()
        context['results'] = list(result)
        return render_template('lookup_execution.html', **context)

@app.route("/kb/<id>", methods=['POST', 'GET'])
def kb_detail(id):

    q1 = ("SELECT * FROM known_knowledge_base k WHERE k.issue = %s;")
    q2 = "SELECT * FROM known_knowledge_base k INNER JOIN scripts s ON k.issue = %s AND k.issue = s.issue_id;"
    q3 = "SELECT DISTINCT NE_VERSION FROM network_topology;"
    q4 = "SELECT DISTINCT VISIBILITY FROM known_knowledge_base;"
    q5 = "SELECT DISTINCT NE_TYPE FROM network_topology;"
    q6 = "SELECT DISTINCT action_type FROM actions;"

    res1 = db.execute_query(q1, (id,), fetch_all=True)
    res2 = db.execute_query(q2, (id,), fetch_all=True)
    res3 = db.execute_query(q3, fetch_all=True)
    res4 = db.execute_query(q4, fetch_all=True)
    res5 = db.execute_query(q5, fetch_all=True)
    res6 = db.execute_query(q6, fetch_all=True)

    _len_script = len(res2)
    return render_template('DetailKbPage.html', records=res1, users=res2, software=[s['NE_VERSION'] for s in res3],
                           visibilities=[i['VISIBILITY'] for i in res4], technologies=[i['NE_TYPE'] for i in res5],
                           action_type=[i['action_type'] for i in res6],len_script=_len_script)


@app.route("/save_singleissue", methods=['POST', 'GET'])
def save_singleissue():
    issue = request.form.get('msId')
    poc = request.form.get('POC')
    sw = request.form.get('swname')
    status = request.form.get('status')
    area = request.form.get('area')
    tech = request.form.get('tech')
    visible = request.form.get('Visibility')
    kpi = request.form.get('kpi')
    detect = request.form.get('detect')
    cell = request.form.get('cell')
    desc = request.form.get('desc')
    rca = request.form.get('rca')
    impact = request.form.get('impact')
    script = request.form.get('Detection_Methodology_Log_Script')
    oms = request.form.get('Detection_Methodology_OMs')
    wa = request.form.get('Interim_Solution_Workaround')
    pf = request.form.get('Permanent_Fix')
    ic = request.form.get('ic')
    print("status",status)
    update_query = "UPDATE known_knowledge_base SET Poc=%s, SW=%s, Status=%s, Area=%s, Technology=%s, Major_KPI_Degradation=%s, Issue_Category=%s, " \
                   "Description=%s, Root_Cause=%s, Impact=%s, Detection_Methodology_OMs=%s, Detection_Methodology_Log_Script=%s, " \
                   "Interim_Solution_Workaround=%s,Visibility=%s, Permanent_Fix=%s, Issue_Detection_Category=%s, Cell_or_Site_Level=%s WHERE issue=%s"
    values = (poc, sw, status, area, tech, kpi, ic, desc, rca, impact, oms, script, wa, visible, pf, detect, cell, issue)
    try:
        db.execute_query(update_query, values, commit=True)
        return jsonify({'message': "updated succesfully!! Click on next to update action steps"})
    except Exception as e:
        return jsonify({'error': f"Something went wrong {e}"})

@app.route('/execute-multi', methods=['POST','GET'])
def execute_multi():
    
    data = request.get_json()
    issues = data.get('issues', [])

    results = []
    script_args_list = []
    exec_start = datetime.datetime.now()
    exec_end = exec_start + datetime.timedelta(minutes=2)

    for issue_data in issues:
        dbval = db.execute_query("select * from known_knowledge_base where issue=%s", (issue_data['issue'],))
        om_check = dbval[0]['#1_OM_validation_check']
        om_cnt = dbval[0]['OM_Family_OM_Counters']
        sites = issue_data['sites']
        issue = issue_data['issue']
        session_id = ifast_rand(8, issue_data['issue'])
        results.append({
            'session_id': session_id,
            'user_name': session['user_name'],
            'issue': issue,
            'issue_category': issue_data['issue_category'],
            'curr_action_no': '3/3',
            'timestamp_start': exec_start.strftime("%Y-%m-%d %H:%M:%S"),
            'timestamp_end': exec_end.strftime("%Y-%m-%d %H:%M:%S"),
            'ne_matches': '-'
        })
        script_args_list.append([om_cnt, om_check, sites, issue, "ex", "DEV"])

    current_script_args_list = script_args_list
    multiprocessing.Process(target=execute_scripts, args=(script_args_list,)).start()

    return jsonify({'results': results})


def execute_scripts(script_args_list):
    app.logger.info("Starting parallel script execution")
    app.logger.info(f"{script_args_list}")
    queues = [multiprocessing.Queue() for _ in range(len(script_args_list))]
    processes = []

    for i, script_args in enumerate(script_args_list):
        process_id = f"process_{i+1}"
        p = multiprocessing.Process(target=run_scripts, args=(process_id,script_args, queues[i]))
        processes.append(p)
        p.start()

    def generate():
        try:
            while any(p.is_alive() for p in processes):
                for queue in queues:
                    while not queue.empty():
                        message = queue.get()
                        app.logger.info(f"Queue message: {message}")
                        yield f"data: {message}\n\n"
                time.sleep(1)  # Sleep to prevent high CPU usage

            yield "data: [END OF STREAM]\n\n"

        except Exception as e:
            app.logger.error(f"Error occurred: {str(e)}")
            yield "data: [ERROR OCCURRED]\n\n"
        finally:
            app.logger.info("Parallel script execution completed.")

    return Response(stream_with_context(generate()), mimetype='text/event-stream')


def ifast_rand(length,issue):
    random_val = ''.join(random.choices('0123456789', k=length))
    random_val = f'iFAST_{issue}_{random_val}'
    return random_val

@app.route('/status-updates')
def status_updates():
    def generate():
        while True:
            for session_id, status in task_status.items():
                if status in ['Completed', 'Failed']:
                    yield f"data: {{'session_id': '{session_id}', 'status': '{status}'}}\n\n"
                    # Optionally, clear the status after sending update
                    del task_status[session_id]
            time.sleep(5)  # Adjust the polling interval as needed

    return Response(stream_with_context(generate()), mimetype='text/event-stream')

@app.route('/task-status')
def task_status_endpoint():
    return jsonify(task_status)

@app.route('/stream-logs', methods=['GET'])
def stream_logs():
    global current_script_args_list 
    def generate():
        try:
            queues = [multiprocessing.Queue() for _ in range(len(current_script_args_list))]
            processes = []
            for i, script_args in enumerate(current_script_args_list):
                p = multiprocessing.Process(target=run_scripts, args=(script_args, queues[i]))
                processes.append(p)
                p.start()

            while any(p.is_alive() for p in processes):
                for queue in queues:
                    while not queue.empty():
                        message = queue.get()
                        app.logger.info(f"Queue message: {message}")
                        yield f"data: {message}\n\n"
                time.sleep(1)  

            yield "data: [END OF STREAM]\n\n"

        except Exception as e:
            app.logger.error(f"Error occurred: {str(e)}")
            yield "data: [ERROR OCCURRED]\n\n"

    return Response(stream_with_context(generate()), mimetype='text/event-stream')

   


if __name__ == '__main__':
    app.run(port=5004, debug=True)
    socketio.run(app, host='0.0.0.0', port=8000,threaded=True)

